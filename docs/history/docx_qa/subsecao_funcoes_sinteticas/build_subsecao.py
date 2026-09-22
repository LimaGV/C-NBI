from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


REFERENCE = Path(
    r"C:\Users\gabri\Documents\Dissertação\01_TEXTOS\capitulos\dissertação CNBI_16.07.docx"
)
OUTPUT = Path(
    r"C:\Users\gabri\Documents\Dissertação\01_TEXTOS\capitulos\subsecao_metodologica_funcoes_sinteticas.docx"
)
EXPECTED_REFERENCE_SHA256 = (
    "52a1b7c0985a61939f4be31b44cf8b10ac0d3ace4c497b16c680d5201727b172"
)


TITLE = "Construção dos cenários e das funções sintéticas"

PARAGRAPHS = [
    (
        "Os cenários sintéticos foram construídos para avaliar o C-NBI em condições "
        "controladas, nas quais a dificuldade do problema pudesse ser modificada sem "
        "alterar simultaneamente todos os demais elementos do experimento. Essa escolha "
        "permite saber, com maior segurança, se uma mudança de desempenho decorre do "
        "número de objetivos, do grau de dependência entre eles ou do próprio método de "
        "otimização. Foram mantidas três variáveis de decisão em todos os casos e foram "
        "combinados três números de objetivos — 4, 6 e 12 — com três níveis de correlação "
        "estrutural — baixa, média e alta. O cruzamento desses fatores originou nove "
        "cenários distintos."
    ),
    (
        "As três variáveis de decisão foram tratadas em escala codificada. Nessa escala, "
        "o centro representa a condição experimental de referência, enquanto valores "
        "positivos e negativos indicam deslocamentos em direções opostas. A região "
        "admissível para a otimização foi delimitada por uma esfera centrada na origem, "
        "com raio aproximadamente igual a 1,682, além dos limites individuais das "
        "variáveis. Essa região é a mesma para todos os cenários. Desse modo, as "
        "comparações não são confundidas por mudanças no tamanho ou na forma do espaço de "
        "busca."
    ),
    (
        "Cada função objetivo foi definida a partir de um ponto de referência próprio, "
        "denominado âncora. A resposta correspondente mede, em escala quadrática, o "
        "afastamento entre uma condição candidata e sua âncora. Em termos intuitivos, a "
        "superfície assume a forma de uma bacia suave: o menor valor ocorre exatamente na "
        "âncora e a resposta cresce progressivamente à medida que a condição se afasta "
        "dela. Como são usadas distâncias ao quadrado, todas as respostas são contínuas, "
        "convexas, não negativas e possuem derivadas bem definidas. Essas propriedades "
        "tornam o problema numericamente controlável sem torná-lo trivial para a "
        "otimização multiobjetivo."
    ),
    (
        "O conflito entre os objetivos surge porque suas âncoras ocupam posições "
        "diferentes. Aproximar-se da melhor condição de uma resposta geralmente implica "
        "afastar-se da melhor condição de outras respostas. Portanto, o compromisso não é "
        "introduzido artificialmente por sinais trocados ou por restrições específicas de "
        "cada objetivo; ele decorre diretamente da geometria comum das superfícies. Essa "
        "construção também oferece uma interpretação clara para os extremos do problema: "
        "cada âncora é o ótimo individual conhecido de uma resposta."
    ),
    (
        "A dependência entre as respostas também foi controlada pela disposição das "
        "âncoras. Para isso, avaliou-se cada conjunto candidato de âncoras em uma amostra "
        "de condições distribuídas de maneira uniforme pelo espaço admissível. Em seguida, "
        "calculou-se a correlação linear de Pearson entre cada par de respostas e adotou-se "
        "a média dos valores absolutos como medida-resumo do cenário. Assim, uma "
        "correlação estrutural baixa representa respostas que variam de maneira mais "
        "distinta ao longo do espaço de decisão; uma correlação alta representa respostas "
        "que tendem a variar de forma mais semelhante. É importante destacar que essa "
        "dependência é produzida pela geometria das funções verdadeiras, e não por erros "
        "experimentais correlacionados."
    ),
    (
        "As posições das âncoras foram determinadas por uma busca numérica determinística "
        "com múltiplos pontos iniciais. A busca ajustou simultaneamente a direção e a "
        "distância de cada âncora em relação ao centro, procurando atingir correlações "
        "médias de 0,25, 0,60 e 0,85. Além da proximidade com o nível desejado, foram "
        "impostas condições para impedir âncoras coincidentes, excessivamente próximas ou "
        "incapazes de ocupar as três direções do espaço de decisão. A confirmação final "
        "foi realizada com 50.000 pontos de uma sequência de Sobol, que oferece cobertura "
        "mais uniforme do espaço do que uma amostra aleatória simples. Um cenário somente "
        "foi aceito quando a correlação obtida diferiu do alvo em, no máximo, 0,03."
    ),
    (
        "A utilização de 6 ou 12 objetivos com apenas três variáveis de decisão cria "
        "redundância de forma intencional. Todas as respostas são construídas a partir das "
        "mesmas três coordenadas e compartilham um componente comum de curvatura. Por isso, "
        "o acréscimo de objetivos não acrescenta, na mesma proporção, novas direções "
        "geométricas independentes. Essa característica reproduz a situação de interesse "
        "da pesquisa: problemas com muitas respostas observadas, mas com uma estrutura "
        "latente de dimensão menor, nos quais a aplicação direta do NBI pode se tornar "
        "dimensionalmente incompatível."
    ),
    (
        "Depois de definidas as funções verdadeiras, foi simulado o processo de obtenção "
        "dos dados experimentais. Empregou-se um planejamento composto central com 19 "
        "ensaios: oito combinações fatoriais, seis pontos axiais e cinco repetições no "
        "centro. As respostas verdadeiras foram calculadas nesses pontos e receberam erros "
        "gaussianos independentes. A intensidade do erro foi calibrada separadamente para "
        "cada objetivo, de forma que o sinal representasse aproximadamente 95% da "
        "variabilidade esperada. Com isso, todos os cenários apresentam elevado poder "
        "explicativo, mas preservam a incerteza típica de dados experimentais."
    ),
    (
        "Para cada objetivo foi então ajustado um modelo completo de superfície de resposta "
        "de segunda ordem pelo método dos mínimos quadrados ordinários. O modelo inclui um "
        "termo constante, os efeitos individuais das três variáveis, suas curvaturas e as "
        "interações entre pares de variáveis. Essa estrutura é suficientemente flexível "
        "para representar exatamente as funções sintéticas na ausência de ruído. Com "
        "ruído, o ajuste passa a representar aquilo que estaria disponível em uma aplicação "
        "experimental real. Os métodos de otimização receberam exclusivamente esses "
        "modelos ajustados; as funções verdadeiras não foram consultadas durante a busca."
    ),
    (
        "A separação entre função verdadeira e modelo ajustado é central para a validade do "
        "experimento. A função verdadeira representa o processo conhecido pelo pesquisador "
        "que construiu o benchmark e é usada somente para gerar observações, estabelecer a "
        "fronteira de referência e avaliar os resultados ao final. O modelo de superfície "
        "de resposta representa o conhecimento efetivamente disponível aos algoritmos. "
        "Além disso, em cada repetição, todos os métodos comparados utilizaram exatamente o "
        "mesmo conjunto de modelos ajustados. Essa estratégia pareada evita que diferenças "
        "entre amostras de ruído sejam confundidas com diferenças entre métodos."
    ),
    (
        "A construção adotada também permite conhecer a região de Pareto verdadeira sem "
        "recorrer a um algoritmo evolucionário. Para essa classe de funções, as soluções "
        "eficientes encontram-se no casco convexo formado pelas âncoras, isto é, na menor "
        "região que contém todas elas e todos os seus compromissos convexos. A referência "
        "foi formada por 100.000 pontos. Primeiro, todas as âncoras foram incluídas "
        "explicitamente. Depois, o casco foi dividido em tetraedros e os demais pontos "
        "foram distribuídos entre eles de acordo com seus volumes. Dentro de cada "
        "tetraedro, os pontos foram obtidos por combinações convexas aleatórias. Esse "
        "procedimento produz uma referência reprodutível e independente dos métodos "
        "avaliados."
    ),
    (
        "Por fim, a geração dos cenários foi acompanhada por verificações numéricas. Foi "
        "confirmado que todas as âncoras são distintas, pertencem à região interna e ocupam "
        "as três dimensões do espaço de decisão; que cada âncora realmente fornece o mínimo "
        "conhecido de sua resposta; que os níveis de correlação respeitam a tolerância "
        "estabelecida; e que, sem ruído, o modelo quadrático recupera as funções verdadeiras "
        "com erro numérico desprezível. Também se verificou que a referência de Pareto é "
        "reprodutível, permanece no casco das âncoras e apresenta, em amostras de auditoria, "
        "pelo menos 99% de pontos não dominados. Em conjunto, essas verificações asseguram "
        "que os nove cenários diferem nos fatores planejados, mas compartilham uma base "
        "geométrica, estatística e computacional comum."
    ),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_font(run, *, bold: bool = False, size: float = 12.0) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "pt-BR")


def set_outline_level(style, level: int) -> None:
    ppr = style.element.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), str(level))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run, size=10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def build() -> None:
    if sha256(REFERENCE) != EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("A referência foi alterada desde a destilação do estilo.")

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    styles = document.styles
    if "Corpo do texto Dissertação" in styles:
        body_style = styles["Corpo do texto Dissertação"]
    else:
        body_style = styles.add_style("Corpo do texto Dissertação", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = normal
    body_style.font.name = "Times New Roman"
    body_style.font.size = Pt(12)
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_style.paragraph_format.first_line_indent = Cm(1.0)
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.widow_control = True

    title_style = styles["Heading 2"]
    title_style.base_style = normal
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(12)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    title_style.paragraph_format.first_line_indent = Cm(0)
    title_style.paragraph_format.left_indent = Cm(0)
    title_style.paragraph_format.line_spacing = 1.5
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(6)
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.keep_together = True
    set_outline_level(title_style, 1)

    title = document.add_paragraph(style=title_style)
    title.paragraph_format.first_line_indent = Cm(0)
    set_font(title.add_run(TITLE), bold=True)

    for text in PARAGRAPHS:
        paragraph = document.add_paragraph(style=body_style)
        paragraph.paragraph_format.widow_control = True
        set_font(paragraph.add_run(text))

    header = section.header
    header.paragraphs[0].text = ""
    footer = section.footer
    footer.paragraphs[0].clear()
    add_page_number(footer.paragraphs[0])

    core = document.core_properties
    core.title = TITLE
    core.subject = "Subseção metodológica sobre a construção das funções sintéticas"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = "C-NBI; funções sintéticas; metodologia; superfície de resposta"
    core.comments = "Documento autônomo elaborado a partir da metodologia implementada no repositório CNBI-Synthetic-Benchmarks."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    if sha256(REFERENCE) != EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("A referência foi modificada durante a criação.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
