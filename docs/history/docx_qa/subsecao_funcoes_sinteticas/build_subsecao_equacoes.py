from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

import build_subsecao as base


OUTPUT = Path(
    r"C:\Users\gabri\Documents\Dissertação\01_TEXTOS\capitulos\subsecao_metodologica_funcoes_sinteticas_com_equacoes.docx"
)
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

SCENARIOS = [
    ("m4_low", 4, "Baixa", "0,25", "0,2509"),
    ("m4_medium", 4, "Média", "0,60", "0,5993"),
    ("m4_high", 4, "Alta", "0,85", "0,8506"),
    ("m6_low", 6, "Baixa", "0,25", "0,2501"),
    ("m6_medium", 6, "Média", "0,60", "0,5990"),
    ("m6_high", 6, "Alta", "0,85", "0,8495"),
    ("m12_low", 12, "Baixa", "0,25", "0,2531"),
    ("m12_medium", 12, "Média", "0,60", "0,5998"),
    ("m12_high", 12, "Alta", "0,85", "0,8392"),
]

EQUATIONS = {
    "domain": """
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <mi mathvariant="normal">Ω</mi><mo>=</mo><mo>{</mo>
    <mi mathvariant="bold">x</mi><mo>∈</mo><msup><mi mathvariant="double-struck">R</mi><mn>3</mn></msup>
    <mo>:</mo><msup><mi mathvariant="bold">x</mi><mi mathvariant="normal">T</mi></msup>
    <mi mathvariant="bold">x</mi><mo>≤</mo><msup><mi>α</mi><mn>2</mn></msup><mo>}</mo>
    <mo>,</mo><mspace width="1em"/><mi>α</mi><mo>=</mo><msup><mn>2</mn><mfrac><mn>3</mn><mn>4</mn></mfrac></msup>
    <mo>≈</mo><mn>1,682</mn>
  </mrow>
</math>
""",
    "truth": """
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mi>f</mi><mi>j</mi></msub><mfenced><mi mathvariant="bold">x</mi></mfenced><mo>=</mo>
    <msubsup>
      <mfenced open="‖" close="‖"><mrow><mi mathvariant="bold">x</mi><mo>−</mo><msub><mi mathvariant="bold">a</mi><mi>j</mi></msub></mrow></mfenced>
      <mn>2</mn><mn>2</mn>
    </msubsup>
    <mo>,</mo><mspace width="1em"/><mi>j</mi><mo>=</mo><mn>1</mn><mo>,</mo><mo>…</mo><mo>,</mo><mi>m</mi>
  </mrow>
</math>
""",
    "correlation": """
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <mover><mi>ρ</mi><mo>¯</mo></mover><mo>=</mo>
    <mfrac><mn>2</mn><mrow><mi>m</mi><mfenced><mrow><mi>m</mi><mo>−</mo><mn>1</mn></mrow></mfenced></mrow></mfrac>
    <munder><mo>∑</mo><mrow><mi>i</mi><mo>&lt;</mo><mi>j</mi></mrow></munder>
    <mfenced open="|" close="|"><msub><mi>ρ</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub></mfenced>
  </mrow>
</math>
""",
    "noise": """
<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">
  <mrow>
    <msub><mi>y</mi><mi>j</mi></msub><mfenced><mi mathvariant="bold">x</mi></mfenced><mo>=</mo>
    <msub><mi>f</mi><mi>j</mi></msub><mfenced><mi mathvariant="bold">x</mi></mfenced><mo>+</mo><msub><mi>ε</mi><mi>j</mi></msub>
    <mo>,</mo><mspace width="1em"/><msub><mi>ε</mi><mi>j</mi></msub><mo>∼</mo>
    <mi mathvariant="script">N</mi><mfenced><mrow><mn>0</mn><mo>,</mo><msubsup><mi>σ</mi><mi>j</mi><mn>2</mn></msubsup></mrow></mfenced>
  </mrow>
</math>
""",
}


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_borders(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in edges.items():
        tag = "start" if edge == "left" else "end" if edge == "right" else edge
        element = borders.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            borders.append(element)
        element.set(qn("w:val"), spec.get("val", "single"))
        element.set(qn("w:sz"), str(spec.get("sz", 8)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), spec.get("color", "000000"))


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(header)


def add_equation(document: Document, mathml: str) -> None:
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    math_tree = etree.fromstring(mathml.encode("utf-8"))
    omml = transform(math_tree).getroot()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    paragraph._p.append(parse_xml(etree.tostring(omml)))


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Corpo do texto Dissertação")
    paragraph.paragraph_format.widow_control = True
    base.set_font(paragraph.add_run(text))


def add_small_paragraph(document: Document, text: str, *, align, keep_next=False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(text)
    base.set_font(run, size=10)


def add_scenario_table(document: Document) -> None:
    add_small_paragraph(
        document,
        "Tabela 1 - Estrutura e correlações dos nove cenários sintéticos.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        keep_next=True,
    )
    headers = ["Cenário", "Nº de objetivos", "Nível", "Correlação-alvo", "Correlação obtida"]
    table = document.add_table(rows=1, cols=len(headers))
    widths = [2200, 1600, 1350, 1750, 2050]
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for col, text in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(text)
        base.set_font(run, bold=True, size=10)
        set_cell_borders(cell, top={"sz": 10}, bottom={"sz": 8})
    for scenario, objectives, level, target, achieved in SCENARIOS:
        row = table.add_row()
        values = [scenario, str(objectives), level, target, achieved]
        for col, (cell, text) in enumerate(zip(row.cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(text)
            base.set_font(run, size=10)
    set_table_geometry(table, widths)
    for cell in table.rows[-1].cells:
        set_cell_borders(cell, bottom={"sz": 10})
    add_small_paragraph(
        document,
        "Fonte: Elaborado pelo autor (2026).",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_small_paragraph(
        document,
        "Nota: todos os valores obtidos permaneceram dentro da tolerância absoluta de 0,03 em relação ao alvo.",
        align=WD_ALIGN_PARAGRAPH.LEFT,
    )


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = False

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    body_style = document.styles.add_style("Corpo do texto Dissertação", 1)
    body_style.base_style = normal
    body_style.font.name = "Times New Roman"
    body_style.font.size = Pt(12)
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_style.paragraph_format.first_line_indent = Cm(1.0)
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.widow_control = True

    title_style = document.styles["Heading 2"]
    title_style.base_style = normal
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(12)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    title_style.paragraph_format.first_line_indent = Cm(0)
    title_style.paragraph_format.left_indent = Cm(0)
    title_style.paragraph_format.line_spacing = 1.5
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(6)
    title_style.paragraph_format.keep_with_next = True
    title_style.paragraph_format.keep_together = True
    base.set_outline_level(title_style, 1)


def build() -> None:
    if base.sha256(base.REFERENCE) != base.EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("A referência de estilo foi alterada desde a destilação.")
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Heading 2")
    title.paragraph_format.first_line_indent = Cm(0)
    base.set_font(title.add_run(base.TITLE), bold=True)

    paragraphs = list(base.PARAGRAPHS)
    paragraphs[9] = paragraphs[9].replace("benchmark", "conjunto de testes")
    paragraphs[1] += " Em notação compacta, a região é descrita pela expressão a seguir."
    paragraphs[2] += " A definição matemática utilizada é apresentada em seguida."
    paragraphs[4] += " A medida-resumo empregada na calibração é apresentada a seguir."
    paragraphs[7] += " A observação simulada é representada pela expressão a seguir."

    add_body(document, paragraphs[0])
    add_scenario_table(document)

    add_body(document, paragraphs[1])
    add_equation(document, EQUATIONS["domain"])
    add_body(
        document,
        "Na expressão, o conjunto indicado pela letra grega ômega representa a região admissível, o vetor em negrito reúne as três variáveis codificadas e alfa corresponde ao raio da esfera.",
    )

    add_body(document, paragraphs[2])
    add_equation(document, EQUATIONS["truth"])
    add_body(
        document,
        "Nessa expressão, a função com índice j representa a j-ésima resposta, o vetor de decisão reúne as três condições codificadas, a âncora com o mesmo índice define seu ótimo individual e m é o número total de objetivos. O símbolo de norma representa a distância euclidiana.",
    )

    add_body(document, paragraphs[3])
    add_body(document, paragraphs[4])
    add_equation(document, EQUATIONS["correlation"])
    add_body(
        document,
        "O termo à esquerda representa a correlação estrutural média. Cada parcela no somatório corresponde ao valor absoluto da correlação de Pearson entre um par distinto de respostas; o fator inicial transforma a soma na média de todos os pares possíveis.",
    )

    add_body(document, paragraphs[5])
    add_body(document, paragraphs[6])
    add_body(document, paragraphs[7])
    add_equation(document, EQUATIONS["noise"])
    add_body(
        document,
        "A resposta observada é, portanto, a soma da função verdadeira e de um erro aleatório específico do objetivo. Os erros foram independentes, com média zero e variância calibrada separadamente para manter aproximadamente 95% da variabilidade associada ao sinal.",
    )

    for text in paragraphs[8:]:
        add_body(document, text)

    section = document.sections[0]
    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].clear()
    base.add_page_number(section.footer.paragraphs[0])

    core = document.core_properties
    core.title = base.TITLE
    core.subject = "Subseção metodológica com equações e tabela dos cenários sintéticos"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = "C-NBI; funções sintéticas; metodologia; superfície de resposta"
    core.comments = "Versão revisada com quatro equações e tabela dos nove cenários."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    if base.sha256(base.REFERENCE) != base.EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("A referência foi modificada durante a edição.")
    print(OUTPUT)


if __name__ == "__main__":
    build()
