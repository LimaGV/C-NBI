from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX = Path(__file__).resolve().parents[2] / 'Papper/Manuscript - CNBI.docx'
OUT = DOCX.with_name('Manuscript - CNBI - Section 2.docx')
BACKUP = DOCX.with_name('Manuscript - CNBI.before-section2.docx')
if not BACKUP.exists():
    BACKUP.write_bytes(DOCX.read_bytes())

doc = Document(DOCX)
styles = {s.name for s in doc.styles}
body_style = 'Normal (Web)' if 'Normal (Web)' in styles else 'Normal'

# Locate the placeholder and the first declaration that follows it.
section_heading = next(p for p in doc.paragraphs if p.text.strip() == 'Theoretical background')
funding = next(p for p in doc.paragraphs if p.text.strip() == 'Funding')
section_heading.text = 'Combinatorial normal boundary intersection'
section_heading.style = doc.styles['Heading 1']

created = []

def track(element):
    created.append(element)
    return element

def add_heading(text, level=2):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 2 else 7)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    track(p._p)
    return p

def add_body(text, first=True):
    # Keep mathematical symbols readable in running text without leaving
    # programming-style underscores in the manuscript.
    replacements = {
        'x_j^*': 'xⱼ*', 'x_i^*': 'xᵢ*',
        'f_i^U': 'fᵢᵁ', 'f_j': 'fⱼ', 'f_i': 'fᵢ',
        'a_i': 'aᵢ', 'a_j': 'aⱼ',
        'n_x': 'nₓ', 'q_S': 'qₛ', 'E_S': 'Eₛ', 'E_G': 'Eɢ',
        'A_S': 'Aₛ', 'n_S': 'nₛ', 'B_k': 'Bₖ', 'H_k': 'Hₖ',
        'δ_k': 'δₖ', 's_ij': 'sᵢⱼ', 'τ_r': 'τᵣ',
        'MSE_j': 'MSEⱼ',
        'a_1^T': 'a₁ᵀ', 'a_k^T': 'aₖᵀ',
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    p = doc.add_paragraph(style=body_style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(5)
    if first:
        p.paragraph_format.first_line_indent = Inches(0.25)
    p.add_run(text)
    track(p._p)
    return p

def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for key, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+key))
        if node is None:
            node = OxmlElement('w:'+key); tcMar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'),'dxa')

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node = borders.find(qn('w:'+edge))
        if node is None:
            node = OxmlElement('w:'+edge); borders.append(node)
        node.set(qn('w:val'),'nil')

def add_equation(linear, number):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    widths = [0.20, 6.25, 0.80]
    for idx, (cell,w) in enumerate(zip(table.rows[0].cells,widths)):
        cell.width = Inches(w); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx == 2:
            set_cell_margins(cell, start=20, end=20)
        else:
            set_cell_margins(cell)
    mid = table.cell(0,1).paragraphs[0]
    mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mid.paragraph_format.space_before = Pt(3); mid.paragraph_format.space_after = Pt(3)
    run = mid.add_run(linear); run.font.name='Cambria Math'; run.font.size=Pt(10.5)
    right = table.cell(0,2).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_before = Pt(3); right.paragraph_format.space_after = Pt(3)
    rr=right.add_run(f'({number})'); rr.font.name='Times New Roman'; rr.font.size=Pt(10)
    track(table._tbl)
    return table

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def add_algorithm():
    cap=doc.add_paragraph(style=body_style)
    cap.paragraph_format.keep_with_next=True; cap.paragraph_format.space_before=Pt(8); cap.paragraph_format.space_after=Pt(3)
    r=cap.add_run('Algorithm 1. Combinatorial normal boundary intersection'); r.bold=True
    track(cap._p)
    steps=[
      ('1','Orient all objectives to minimization and determine each individual optimum.'),
      ('2','Construct and normalize the global payoff matrix and its anchor decisions.'),
      ('3','Estimate the effective geometric dimension and the global spectral window.'),
      ('4','Enumerate every objective subset with 2 ≤ k ≤ min(m, nₓ + 1).'),
      ('5','Build the local edge matrix and reject exactly degenerate subsets.'),
      ('6','Retain a subset only when its weakest direction and conditioning satisfy the spectral window.'),
      ('7','Generate Simplex-Lattice weights using the resolution assigned to k.'),
      ('8','Solve the local NBI subproblems and retain feasible converged decisions.'),
      ('9','Reevaluate every retained decision in all m original objectives.'),
      ('10','Remove duplicate decisions and apply global nondominance filtering.'),
      ('11','Return the final frontier and an audit ledger linking solutions to subsets and weights.'),
    ]
    table=doc.add_table(rows=1, cols=2)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.columns[0].width=Inches(.55); table.columns[1].width=Inches(6.7)
    hdr=table.rows[0].cells
    hdr[0].text='Step'; hdr[1].text='Operation'
    for c in hdr:
        shade(c,'1F4E78'); set_cell_margins(c,80,90,80,90)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True; run.font.size=Pt(9)
    for idx,(num,text) in enumerate(steps):
        cells=table.add_row().cells; cells[0].text=num; cells[1].text=text
        if idx%2: shade(cells[0],'F3F6F8'); shade(cells[1],'F3F6F8')
        for j,c in enumerate(cells):
            set_cell_margins(c,75,90,75,90); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs: run.font.name='Times New Roman'; run.font.size=Pt(9)
    # light borders
    tblPr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+edge); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); borders.append(node)
    tblPr.append(borders)
    track(table._tbl)

add_body('This section defines the mathematical formulation and computational sequence of combinatorial normal boundary intersection (CNBI). The central distinction is between the objective space used to generate candidates and the objective space used to accept them. Dimensionally admissible subsets support local NBI searches, whereas every resulting decision is evaluated and filtered using the complete set of original objectives. Thus, the decomposition changes the search mechanism without redefining the decision problem.', first=False)

add_heading('2.1 Problem formulation and normalized payoff geometry',2)
add_body('Consider a continuous multiobjective minimization problem with nₓ decision variables and m objective functions. The feasible region Ω includes all variable bounds and problem constraints. Objectives originally defined for maximization are multiplied by −1 before the optimization, so that a common minimization convention is used throughout the method.')
add_equation('min_(x∈Ω) F(x)=[f_1(x),…,f_m(x)]^T,    x∈ℝ^(n_x)',1)
add_body('For each objective f_j, an anchor decision x_j^* is obtained by individual optimization. Evaluating all objectives at every anchor produces the payoff matrix Φ. Its jth column is the complete objective vector generated by the individual optimum of f_j. Off-diagonal entries are retained because they describe the deterioration of the other objectives when f_j is optimized alone (Das & Dennis, 1998).')
add_equation('x_j^*∈arg min_(x∈Ω) f_j(x),    Φ_(ij)=f_i(x_j^*),    i,j=1,…,m',2)
add_body('Objective scales can differ by several orders of magnitude. CNBI therefore uses a common affine transformation based on the ideal value f_i^U and the amplitude a_i represented in the payoff geometry. For ordinary RSM problems, f_i^U is the row minimum of Φ and a_i is the difference between its maximum and minimum. Benchmark functions with known theoretical scales use the declared ideal and amplitude. The same transformation is applied to the payoff matrix and to every objective vector used inside an NBI subproblem.')
add_equation('f̄_i(x)=(f_i(x)-f_i^U)/a_i,    a_i=f_i^N-f_i^U>0',3)
add_body('Let Φ̄ contain the normalized payoff columns. Their convex hull of individual minima (CHIM) is parameterized by a vector β in the unit simplex. The weights select geometric points on the CHIM; they are not preference coefficients used to aggregate the objective functions.')
add_equation('CHIM=Φ̄β,    β∈Δ_(m-1);    Δ_(m-1): β∈ℝ^m, β_i≥0, Σ_(i=1)^m β_i=1',4)

add_heading('2.2 Direct NBI and dimensional compatibility',2)
add_body('For a fixed β, direct NBI searches from the corresponding CHIM point in an outward normal direction n. The displacement t is maximized until the normalized attainable objective boundary is reached. The equality form used in this study is given in Eq. (5), with feasibility enforced through x ∈ Ω.')
add_equation('max_(x,t) t : F̄(x)-Φ̄β-t n=0,    x∈Ω',5)
add_body('Equation (5) contains m scalar equalities. Once β is fixed, its unknowns are the nₓ decision variables and the scalar t, totaling nₓ + 1 unknowns. Ignoring special dependencies among equations, a necessary counting condition for the system not to be generically overdetermined is therefore given by Eq. (6).')
add_equation('m≤n_x+1',6)
add_body('When m > nₓ + 1, a direct solution can occur only under particular dependencies in the objective mapping. This dimensional restriction is separate from statistical correlation. Even when Eq. (6) is satisfied, anchor points can be linearly or nearly linearly dependent, causing the CHIM to collapse or become poorly conditioned. CNBI addresses the equation-count restriction through temporary decomposition and addresses weak geometry through spectral screening.')

add_heading('2.3 Temporary combinatorial decomposition',2)
add_body('Let I={1,…,m} denote the complete objective index set. CNBI enumerates all unordered subsets S⊆I whose cardinality k is at least two and does not exceed either m or the direct NBI capacity nₓ + 1. The candidate family is defined by Eq. (7).')
add_equation('C=⋃_(k=2)^(min(m,n_x+1)) \\{S⊆I: |S|=k}',7)
add_body('The number of candidate subsets follows directly from the binomial coefficients. This quantity can grow rapidly with m, which motivates screening before local optimization.')
add_equation('|C|=Σ_(k=2)^(min(m,n_x+1)) C(m,k)',8)
add_body('For each S, CNBI extracts from the global normalized payoff matrix the rows and anchor columns associated with the same original objectives. The resulting k × k payoff submatrix defines a local CHIM. The subset is temporary: it specifies which equalities are active in a local NBI search but does not remove the remaining objectives from the final problem. Since an objective can participate in several subsets, the union of local searches explores different projections of the original trade-off structure. This differs from objective-reduction methods that permanently omit responses or replace them with latent variables (Brockhoff & Zitzler, 2006).')

add_heading('2.4 Spectral representation of payoff geometry',2)
add_body('Dimensional admissibility does not guarantee that a local CHIM has the expected geometric dimension. Let A_S=[a_1^T,…,a_k^T]^T be the normalized anchor matrix for subset S, with one anchor in each row. Taking the first anchor as reference produces the edge matrix E_S.')
add_equation('E_S=[(a_2-a_1)^T; … ; (a_k-a_1)^T]∈ℝ^((k-1)×k)',9)
add_body('The singular value decomposition separates the orthogonal directions of the local geometry and their magnitudes (Golub & Van Loan, 2013). A full-dimensional local simplex requires rank(E_S)=k−1. An exactly deficient matrix is rejected because its normal direction is not unique in the form required by the local NBI construction.')
add_equation('E_S=U_S Σ_S V_S^T,    σ_1(S)≥…≥σ_(k-1)(S)≥0',10)
add_body('The smallest singular value measures the weakest direction of the local CHIM. A small value indicates that the anchors approach a lower-dimensional configuration. Conditioning measures the relative imbalance between the strongest and weakest directions. A large condition number characterizes an elongated geometry in which uniform changes in simplex weights can correspond to very different displacements in objective space.')
add_equation('q_S=σ_1(S)/σ_(k-1)(S)',11)
add_body('Neither the smallest singular value nor q_S has a universal acceptable threshold. CNBI derives both limits from the normalized global payoff geometry, allowing the criterion to adapt to the structural scale of each problem rather than relying on an arbitrary fixed tolerance.')

add_heading('2.5 Effective dimension and noise-aware calibration',2)
add_body('The global edge matrix E_G is constructed from all normalized payoff anchors and decomposed by SVD. In RSM applications, the anchors are themselves obtained from estimated models. Small observed directions can therefore reflect either real trade-off structure or prediction uncertainty. CNBI adapts Horn’s parallel-analysis principle to distinguish these possibilities (Horn, 1965).')
add_body('For objective j at anchor decision x_i^*, the standard error of the estimated mean response is calculated from the residual mean square MSE_j, the model vector z(x_i^*), and the inverse information matrix. Division by a_j places this uncertainty on the same normalized scale as the payoff geometry (Myers et al., 2016).')
add_equation('s_(ij)=√(MSE_j z(x_i^*)^T (X^T X)^(-1) z(x_i^*))/a_j.',12)
add_body('CNBI generates B=2,000 independent Gaussian perturbation matrices with element-specific standard deviations s_ij. Each realization is converted to an edge matrix and decomposed by SVD. For spectral position r, τ_r is the 95th percentile of the corresponding simulated singular values.')
add_equation('ε_(ij)^((b))∼N(0,s_(ij)^2),    τ_r=Q_(0.95)\\{σ_r(E_ε^((b))): b=1,…,2000}',13)
add_body('The effective dimension d is the number of observed global directions that exceed their position-specific noise thresholds. The implementation retains a minimum operational dimension of one when no observed direction exceeds its threshold. This lower bound prevents a zero-dimensional diagnostic; it does not restrict CNBI to subsets with only d objectives.')
add_equation('d=max(1, Σ_r 1[σ_r(E_G)>τ_r])',14)
add_body('The benchmark stress tests do not supply fitted-model prediction variances. For MaF8, MaF9, and MaF13, CNBI therefore uses a separate nonparametric calibration: N=500 reference-front objective vectors are standardized, each objective column is independently permuted B=1,000 times, and the observed eigenvalue spectrum is compared sequentially with the 95th-percentile null spectrum. This permutation procedure estimates d for the benchmark diagnostic only; it does not alter the uncertainty-based RSM procedure used in the synthetic DOE and turning application.')
add_body('The effective dimension is converted into a global spectral window. The first discarded global direction supplies the magnitude floor, while the condition number of the retained global subspace supplies the maximum admissible anisotropy. When every available direction is retained, the floor is zero.')
add_equation('τ_σ=σ_(d+1)(E_G), d<rank(E_G);    τ_σ=0, d=rank(E_G);    τ_q=σ_1(E_G)/σ_d(E_G)',15)

add_heading('2.6 Spectral selection and local NBI solution',2)
add_body('A dimensionally admissible subset proceeds to optimization only if its weakest direction exceeds the global residual floor and its condition number does not exceed that of the retained global subspace. The joint rule in Eq. (16) rejects geometries that are too thin in absolute terms, too anisotropic in relative terms, or exactly rank deficient.')
add_equation('retain S ⇔ rank(E_S)=k-1,    σ_(k-1)(S)>τ_σ,    q_S≤τ_q.',16)
add_body('This window is a heuristic screening rule derived from the problem’s own geometry. It is not claimed as a necessary or sufficient condition for NBI convergence. Its purpose is to avoid spending optimization effort on sub-CHIMs whose normal directions are poorly supported by the global structure or indistinguishable from the uncertainty scale.')
add_body('For every retained subset, weight vectors are generated by a Simplex-Lattice design. Let δ_k denote the lattice resolution and H_k=1/δ_k the number of divisions. The final campaign uses δ_k=0.20 for k≤4 and δ_k=0.50 for k>4. Thus, smaller subsets receive a finer geometric search, whereas higher-dimensional subsets use a coarser lattice to control combinatorial growth.')
add_equation('B_k: β∈Δ_(k-1),    H_k β_i∈ℕ_0,    H_k=1/δ_k.',17)
add_body('The number of local NBI subproblems contributed by one retained subset is given by Eq. (18). Consequently, the spectral filter affects cost twice: it removes entire subsets and all lattice points associated with them.')
add_equation('|B_k|=C(H_k+k-1,k-1)',18)
add_body('For a retained subset S and weight β∈B_k, the normalized local NBI problem uses the local anchor matrix A_S and its outward normal n_S. The implementation minimizes −t, which is equivalent to maximizing t in Eq. (19).')
add_equation('max_(x,t) t : F̄_S(x)-β^T A_S-t n_S=0,    x∈Ω',19)
add_body('The local problems are solved by sequential least-squares programming. For quadratic RSM objectives, analytical objective gradients and equality Jacobians are used. For general benchmark functions, finite-difference Jacobians are fully counted as objective-vector evaluations. Exact anchor weights reuse the corresponding individual-optimum decisions. Interior weights use the preceding accepted solution, the convex combination of anchor decisions, anchor decisions ordered by weight, the domain center, and reproducibly generated rescue points as candidate starts. A subproblem is accepted only when the decision is feasible and the infinity norm of the NBI equality residual does not exceed 10⁻⁵. The final solver configuration allows 100 iterations per start and uses an internal optimization tolerance of 10⁻¹⁰.')

add_heading('2.7 Recomposition in the complete objective space',2)
add_body('All accepted local decisions are pooled after the subset searches. Each decision is then reevaluated using all m original objective functions supplied to the algorithm, regardless of the subset that generated it. This step restores the complete decision problem before any final quality judgment is made.')
add_body('Post-processing first removes infeasible decisions. Decisions whose coordinates coincide after quantization at 10⁻⁵ are treated as duplicates; within each group, the candidate with the smallest sum of normalized minimization objectives is retained. Global Pareto filtering is then applied in the complete m-dimensional objective space with a numerical dominance tolerance of 10⁻¹⁰. A point generated by a locally efficient subset can therefore be removed when another candidate dominates it after all objectives are considered.')
add_equation('P_CNBI=ND_m(F(x): x∈⋃_(S∈C_retain) X_S)',20)
add_body('CNBI returns the complete nondominated set as its primary output. It also records the generating subset, lattice weight, convergence status, equality residual, feasibility result, and number of attempts for each candidate. This audit trail makes it possible to distinguish screening decisions, numerical failures, duplicate generation, and removal by global dominance. Algorithm 1 summarizes the complete procedure.')
add_algorithm()
add_body('The worst-case number of local subproblems before screening is the sum, over all admissible cardinalities, of the number of objective subsets multiplied by the corresponding lattice size. The realized cost is lower when the spectral window rejects combinations. Because each nonlinear subproblem can require a different number of objective and gradient evaluations, the experimental comparison reports measured objective-vector evaluations rather than relying only on the combinatorial upper bound.', first=False)

# Move every newly appended XML element immediately before Funding.
for element in created:
    funding._p.addprevious(element)

# Avoid a title stranded from its first paragraph/table.
for p in doc.paragraphs:
    if p.text.strip().startswith('2.'):
        p.paragraph_format.keep_with_next = True

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print('backup', BACKUP)
print('equations', 20)
