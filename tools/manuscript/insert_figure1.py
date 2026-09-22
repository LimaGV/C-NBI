from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = Path(__file__).resolve().parents[2] / 'Papper/Manuscript - CNBI - Section 2.docx'
out = src.with_name('Manuscript - CNBI - Section 2 with Figure.docx')
fig = src.with_name('Figure 1 - CNBI geometric motivation.png')
doc = Document(src)
anchor = next(p for p in doc.paragraphs if p.text.startswith('Neither the smallest singular value'))
anchor.add_run(' The contrast is illustrated in Fig. 1: an additive shift produces no new trade-off direction, whereas genuinely conflicting objectives define a nondegenerate local geometry.')

pfig = doc.add_paragraph()
pfig.alignment = WD_ALIGN_PARAGRAPH.CENTER
pfig.paragraph_format.space_before = Pt(5)
pfig.paragraph_format.space_after = Pt(3)
pfig.paragraph_format.keep_with_next = True
pfig.add_run().add_picture(str(fig), width=Inches(6.45))

cap_style = 'Caption' if 'Caption' in [s.name for s in doc.styles] else 'Normal'
cap = doc.add_paragraph(style=cap_style)
cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
cap.paragraph_format.space_after = Pt(7)
cap.paragraph_format.keep_with_next = True
r = cap.add_run('Fig. 1. Geometric motivation for the spectral screening used by CNBI. ')
r.bold = True
r.font.size = Pt(9)
r2 = cap.add_run('(a) Objectives related by an additive constant generate redundant payoff directions. (b) Conflicting objectives define a nondegenerate CHIM direction (blue), a Pareto-front segment (green), and the NBI displacement t (orange).')
r2.bold = False
r2.font.size = Pt(9)

anchor._p.addnext(pfig._p)
pfig._p.addnext(cap._p)
doc.save(out)
print(out)
