from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(__file__).resolve().parents[2] / 'manuscript/CNBI_ESWA_Manuscript_Draft.docx'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph(style="Body Text")
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.25)
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(caption)
        r.bold = True
        r.font.size = Pt(9.5)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = str(text)
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_border(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(8.5)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            set_cell_border(cells[i])
            if ridx % 2:
                set_cell_shading(cells[i], "F3F6F8")
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    # Keep compact result tables together so Word does not strand a repeated
    # header at the bottom of a page.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(5)

styles["Body Text"].font.name = "Times New Roman"
styles["Body Text"].font.size = Pt(10.5)
styles["Body Text"].font.color.rgb = RGBColor(0, 0, 0)
styles["Body Text"].paragraph_format.line_spacing = 1.12
styles["Body Text"].paragraph_format.space_after = Pt(5)

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style.font.size = Pt(17)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(13)
# Word's built-in Title style can inherit a colored bottom border from the theme.
# Remove it so the manuscript title remains plain, as required for journal copy.
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size, before, after in (("Heading 1", 13, 13, 5), ("Heading 2", 11.5, 10, 4), ("Heading 3", 10.5, 8, 3)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("CNBI manuscript draft  |  ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(footer)

# Front matter
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Combinatorial normal boundary intersection with spectral screening for structurally dependent many-objective optimization")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Gabriel Victor de Lima · Mirelli de Castro Cesário · Matheus Costa Pereira · Anderson Paulo de Paiva")
r.font.name = "Arial"
r.font.size = Pt(10.5)
r.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(11)
r = p.add_run("[Affiliations, ORCID iDs, corresponding author, and e-mail to be confirmed]")
r.italic = True
r.font.name = "Arial"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(100, 100, 100)

add_heading(doc, "Title page information", 1)
p = doc.add_paragraph()
p.add_run("Affiliations: ").bold = True
p.add_run("[Provide the full institutional name and postal address for each affiliation, identified by lower-case superscript letters.]")
p = doc.add_paragraph()
p.add_run("Corresponding author: ").bold = True
p.add_run("[Confirm one corresponding author, full postal address, and active e-mail address.]")
p = doc.add_paragraph()
p.add_run("Acknowledgements: ").bold = True
p.add_run("[Confirm any acknowledgements. This information will remain on the separate title-page file.]")
p = doc.add_paragraph()
p.add_run("Funding: ").bold = True
p.add_run("[Confirm funding agencies, grant numbers, recipients, and sponsor roles, or state that no specific funding was received.]")
p = doc.add_paragraph()
p.add_run("Declaration of competing interests: ").bold = True
p.add_run("The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper. [Confirm with all authors.]")
p = doc.add_paragraph()
p.add_run("Author contributions: ").bold = True
p.add_run("[Complete the CRediT statement after all authors confirm their roles.]")

doc.add_page_break()

add_heading(doc, "Abstract", 1)
abstract = (
    "Engineering decision problems often involve many dependent objectives, whereas classical normal boundary "
    "intersection becomes dimensionally inadmissible when the number of objectives exceeds the number of decision "
    "variables plus one. This study proposes combinatorial normal boundary intersection (CNBI), which diagnoses the "
    "effective objective structure, screens geometrically informative objective subsets, solves admissible lower-dimensional "
    "normal boundary intersection subproblems, and recomposes all candidates in the original objective space. A controlled "
    "3 × 3 × 3 synthetic experiment varied decision-space dimension, dimensional excess, and structural dependence across "
    "27 scenarios with 10 replications. CNBI was compared with rotated-factor normal boundary intersection and two established "
    "evolutionary algorithms. Using the complete fronts produced by each method, CNBI achieved the best inverted generational distance in "
    "26 scenarios and the best hypervolume in 22. Fixed-cardinality curves showed that CNBI retained the best median coverage "
    "from 5 to 20 points. A nine-scenario ablation reduced evaluations by a median of 70.5%, with limited median losses in "
    "coverage and hypervolume. In an eight-response turning application, CNBI obtained an inverted generational distance of "
    "0.0754 and a hypervolume of 0.5023, remaining competitive with the strongest evolutionary comparator. Tests on MaF8 and MaF9 supported generalization, "
    "whereas MaF13 exposed a limitation involving repeated decisions and extreme solutions. The results indicate that spectral "
    "screening can make normal-boundary search practical for structurally dependent many-objective response-surface problems "
    "while preserving the meaning of the original objectives."
)
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.08
p.paragraph_format.space_after = Pt(7)
p.add_run(abstract)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Keywords: ")
r.bold = True
p.add_run("Many-objective optimization; normal boundary intersection; spectral screening; parallel analysis; response surface methodology; manufacturing decision support")

add_heading(doc, "Highlights", 1)
for item in (
    "CNBI extends normal boundary intersection to structurally dependent objectives.",
    "Spectral screening selects admissible subsets without discarding objectives.",
    "CNBI attains the best complete-front IGD in 26 of 27 synthetic scenarios.",
    "Screening cuts median evaluations by 70.5% in a nine-scenario ablation.",
    "An eight-response turning case confirms practical decision-support value.",
):
    add_bullet(doc, item)

doc.add_page_break()

add_heading(doc, "1 Introduction", 1)
add_body(doc, "Manufacturing process optimization increasingly requires simultaneous decisions about quality, productivity, reliability, resource use, and economic performance. These criteria are usually represented by multiple responses that are correlated because they share the same process mechanisms and conflicting because improving one response can impair another (Chiao & Hamada, 2001; Tong et al., 2005). A machining condition that extends tool life, for example, may reduce productivity or increase cost. The resulting task is therefore not to identify a single optimum, but to construct a set of efficient alternatives that makes the trade-offs visible to engineers and decision makers.")
add_body(doc, "Design of experiments and response surface methodology provide a systematic basis for this task. Designed experiments organize data collection, while response surfaces translate the observed effects of controllable factors into explicit mathematical models within a defined experimental region (Myers et al., 2016). Full quadratic models are especially useful because they represent linear effects, curvature, and interactions using a relatively small number of experiments. Once several responses have been modeled, however, their simultaneous optimization requires a method that can explore the Pareto frontier without hiding the meaning of the original performance measures.")
add_body(doc, "Normal boundary intersection (NBI) is attractive for multiresponse optimization because it uses the individual optima to construct the convex hull of individual minima and generates search directions normal to that hull (Das & Dennis, 1998). This geometry can produce a well-distributed frontier and recover solutions in nonconvex regions that may be missed by simple weighted sums. NBI has consequently been integrated with response surface models in environmental and manufacturing applications and has recently been combined with exact and metaheuristic strategies in process planning (Moura et al., 2018; Mechaacha et al., 2026). Its direct formulation nevertheless has a precise dimensional requirement. With m objectives and nₓ decision variables, an NBI subproblem imposes m objective-space equalities on nₓ decision variables plus the scalar displacement t. When m > nₓ + 1, the system is structurally overdetermined. This situation is common in response surface applications, where a small number of controllable factors can generate many measured responses.")
add_body(doc, "The dimensional restriction is different from correlation. It follows from the number of equations and unknowns and can occur even when the objectives are statistically independent. Correlation creates a second difficulty because strongly dependent responses occupy a lower-dimensional effective structure (Jolliffe & Cadima, 2016). In NBI geometry, the individual-optimum anchors may consequently become nearly collinear or coplanar, weakening the convex hull and making the normal directions numerically unstable. A method designed for many-response NBI must therefore address both dimensional admissibility and geometric degeneracy. Merely selecting subsets with no more than nₓ + 1 objectives resolves the first problem, but does not guarantee that the retained anchors contain enough independent structure for a reliable NBI search.")
add_body(doc, "Many-objective evolutionary algorithms approach the problem from a different direction. As objective count increases, the proportion of mutually nondominated candidates rises, reducing the selection pressure supplied by conventional Pareto ranking and making simultaneous convergence and diversity more difficult (Purshouse & Fleming, 2007; Ishibuchi et al., 2008). Reference-point and decomposition-based methods can search the full objective space without imposing the NBI equation-count condition. NSGA-III uses reference points to maintain diversity (Deb & Jain, 2014), whereas MOEA/D decomposes a multiobjective problem into coordinated scalar subproblems (Zhang & Li, 2007). More recent methods adapt dominance relations, reference vectors, and decomposition structures to the shape of the frontier (Bao et al., 2023; Zhang et al., 2023). These developments demonstrate the flexibility of evolutionary search, but they also show that convergence, diversity, reference-set design, and computational cost become increasingly difficult as the number of objectives grows. For response surface problems, an NBI-based alternative remains valuable because it offers an explicit geometric link between individual optima and the generated compromises.")
add_body(doc, "A second line of research reduces the response system before optimization. Principal component and factor methods replace correlated responses with a smaller set of latent dimensions (Jolliffe & Cadima, 2016; Tong et al., 2005). More general objective-reduction methods can also identify and remove responses considered redundant before or during optimization (Sinha et al., 2013). In the turning application used in this study, Pereira et al. (2025) combined Varimax-rotated factor analysis with NBI. The resulting rotated-factor NBI method reconciles the dimensional requirement of NBI with an eight-response problem and produces a compact set of alternatives. This strategy is effective when the latent factors provide a useful summary of the process. Its frontier is nevertheless generated in a transformed space. The physical responses remain recoverable, but they no longer define the search axes directly, which can complicate the interpretation of how a compromise was produced.")
add_body(doc, "These observations expose a methodological gap. Current options either search the complete many-objective space through evolutionary mechanisms, remove objectives judged redundant, or restore NBI compatibility by replacing the original responses with latent variables (Sinha et al., 2013; Pereira et al., 2025). There is limited support for an NBI formulation that temporarily reduces the number of objectives during candidate generation, diagnoses whether each reduced geometry is informative, and still evaluates every accepted solution against all original responses. The research question addressed here is therefore: how can NBI explore structurally dependent many-response problems while preserving the original objective meanings and avoiding dimensionally or geometrically unsuitable subproblems?")
add_body(doc, "This study answers that question with combinatorial normal boundary intersection (CNBI). CNBI enumerates objective subsets with cardinality k satisfying 2 ≤ k ≤ nₓ + 1 and treats each subset as a temporary candidate-generation problem. Before optimization, a spectral screening procedure examines the singular values of the corresponding convex-hull geometry. Its acceptance window is derived from the global payoff structure and from parallel analysis calibrated to response surface prediction uncertainty. Subsets whose anchors are excessively redundant, degenerate, or indistinguishable from model uncertainty are excluded. NBI is then solved for every retained subset. The resulting decisions are reevaluated in all m original objectives, merged, deduplicated, and filtered for global nondominance. Dimensional reduction is therefore used only to generate candidates; final acceptance and interpretation remain in the complete physical response space.")
add_body(doc, "The evaluation was designed to separate controlled evidence, application evidence, and stress testing. A 3 × 3 × 3 synthetic experiment comprises 27 structural scenarios and 10 replications, varying decision-space dimension, dimensional excess, and dependence among objectives. CNBI is compared with rotated-factor NBI, NSGA-III, and MOEA/D. Complete fronts are the primary result because set cardinality is part of the output delivered by each method; fixed-cardinality curves from 2 to 20 points provide a complementary sensitivity analysis. A nine-scenario ablation compares spectral screening with exhaustive solution of all admissible subsets. The practical study uses eight response surface models for hardened AISI H13 steel turning with PCBN 7025 inserts. MaF8, MaF9, and MaF13 from the scalable MaF suite (Cheng et al., 2017) are used only as CNBI stress tests rather than as part of the factorial method comparison.")
add_body(doc, "The main contributions are as follows:", first_line=False)
for item in (
    "A dimensional formulation that identifies precisely when direct NBI becomes overdetermined and separates this condition from correlation-induced geometric degeneracy.",
    "A combinatorial NBI mechanism that preserves the original objectives by restricting dimensional reduction to candidate generation and recomposing every solution in the complete objective space.",
    "A noise-aware spectral screen that uses singular-value geometry and parallel analysis to reject weak objective subsets before optimization, reducing unnecessary subproblems while retaining auditable links between subsets and solutions.",
    "A layered empirical assessment combining a controlled synthetic design, native- and fixed-cardinality comparisons, a nine-scenario ablation, a manufacturing application, and benchmark stress tests.",
):
    add_bullet(doc, item)
add_body(doc, "The experiments show that CNBI achieved the best complete-front inverted generational distance in 26 of 27 synthetic scenarios and the best hypervolume in 22. Spectral screening reduced evaluations by a median of 70.5% across the ablation scenarios, with limited median quality loss. The turning application confirmed competitive coverage, whereas MaF13 exposed a specific limitation involving repeated decisions and extreme solutions. Reporting both the favorable and adverse cases clarifies where the method is useful and where additional safeguards are required.")
add_body(doc, "The remainder of this paper is organized as follows. Section 2 presents the CNBI formulation and spectral screening procedure. Section 3 describes the synthetic design, comparison methods, performance indicators, application, stress tests, and statistical analysis. Section 4 reports the results, Section 5 discusses their implications, Section 6 states the limitations, and Section 7 presents the conclusions.")

add_heading(doc, "2 Combinatorial normal boundary intersection", 1)
add_heading(doc, "2.1 Problem definition", 2)
add_body(doc, "Consider the constrained minimization of an objective vector F(x) = [f₁(x), …, fₘ(x)]ᵀ over a feasible decision region Ω, with x ∈ ℝⁿˣ. The dimensional excess is defined as δ = m − (nₓ + 1). A positive value indicates that the complete objective system exceeds the conventional NBI dimensional condition. Objectives originally expressed as maximization responses are transformed consistently before optimization and returned to their original interpretation for reporting.")

add_heading(doc, "2.2 Response-surface representation and payoff geometry", 2)
add_body(doc, "Each response is represented by a full quadratic response-surface model (Myers et al., 2016). Individual optimization of every response produces the payoff matrix and the corresponding decision vectors. Payoff values are normalized with a common ideal point and response amplitude before spectral analysis. This common scale prevents objectives with larger numerical units from dominating the geometric diagnosis.")

add_heading(doc, "2.3 Noise-aware spectral diagnosis", 2)
add_body(doc, "CNBI forms a difference representation of the normalized payoff geometry and evaluates its singular-value spectrum. Parallel analysis estimates which components exceed the geometry expected from model uncertainty (Horn, 1965). The retained effective dimension defines a spectral window that is later applied to candidate objective subsets. The procedure therefore screens a subset only when its local geometry is too weak or ill-conditioned relative to the global, uncertainty-aware structure.")

add_heading(doc, "2.4 Combinatorial subsets and NBI solution", 2)
add_body(doc, "Candidate subsets contain k = 2, …, min(m, nₓ + 1) original objectives. For each subset, CNBI checks geometric admissibility and spectral consistency. Retained subsets are solved with NBI using simplex-lattice weights, the convex hull of individual minima, and an outward normal direction. The final implementation uses a lattice resolution of 20% for k ≤ 4 and 50% for k > 4. These values apply to all final experiments reported here.")
add_body(doc, "Every accepted decision vector is reevaluated across all m original objectives. Duplicate decisions and dominated solutions are removed only during declared post-processing. The method consequently uses lower-dimensional searches as temporary computational devices rather than as replacements for the full problem. The output also retains a ledger linking each solution to its generating subset, weight, convergence state, and feasibility residual.")

add_heading(doc, "3 Experimental methodology", 1)
add_heading(doc, "3.1 Controlled synthetic experiment", 2)
add_body(doc, "The synthetic design crosses three decision-space dimensions, three levels of dimensional excess, and three levels of structural dependence. This produces 27 structural scenarios. Ten independent noise replications are used per scenario. The true objectives are squared distances to controlled anchors, and the multidimensional generator preserves the logic of the original synthetic module. Anchor calibration uses a tolerance of 0.04 for the target dependence level.")
add_table(doc,
          ["Factor", "Levels", "Purpose"],
          [
              ("Decision dimension nₓ", "2, 3, 5", "Controls the available decision-space geometry"),
              ("Dimensional excess δ", "1, 3, 5", "Controls how far m exceeds nₓ + 1"),
              ("Structural dependence", "Low, medium, high", "Controls shared objective structure"),
              ("Replications", "10 per scenario", "Represents independent response-surface noise"),
          ], widths=[1.6, 1.5, 3.8], caption="Table 1. Controlled synthetic design")

add_heading(doc, "3.2 Compared methods and stopping rules", 2)
add_body(doc, "The comparison includes CNBI, rotated-factor NBI, NSGA-III, and MOEA/D. Rotated-factor NBI retains at least two factors or enough factors to explain 90% of the variation and uses a loading-based factor-score fallback. CNBI and rotated-factor NBI run until their deterministic procedures finish. The evolutionary methods use calibrated population-dependent limits corresponding to approximately 400 complete generations, with evaluation ceilings from 50,000 to 400,400 depending on objective count.")

add_heading(doc, "3.3 Performance indicators and cardinality policy", 2)
add_body(doc, "Inverted generational distance measures convergence and global coverage and is minimized. Generational distance measures the average proximity of produced points to the reference frontier and is minimized. Hypervolume measures dominated objective-space volume and is maximized. All distance metrics use declared objective normalization, and hypervolume uses a common reference point within each configuration.")
add_body(doc, "Complete method outputs constitute the primary comparison because the number of alternatives is part of the service delivered by each method. A complementary sensitivity analysis reduces fronts by hierarchical clustering at fixed cardinalities of 2, 5, 10, and 20 points. A scenario–seed block enters a fixed-cardinality comparison only when every method contains enough points. This rule prevents rich fronts from being compressed silently to the smallest output observed anywhere in the campaign.")

add_heading(doc, "3.4 Additional validation blocks", 2)
add_body(doc, "The MaF benchmark suite (Cheng et al., 2017) is used as a CNBI-only stress test and does not enter the method-comparison statistics or the factorial analysis. MaF8 and MaF9 are evaluated with 4, 6, 8, 10, and 15 objectives; MaF13 is evaluated with 8, 10, and 15 objectives. Each configuration uses 10 seeds and no evaluation ceiling. The turning application contains three decision variables and eight responses describing tool life, reliability, wear, surface quality, process capability, return on investment, and equipment effectiveness. Its data were previously used by Pereira et al. (2025); the present contribution is the CNBI formulation and common comparison protocol.")
add_body(doc, "The ablation compares CNBI with spectral screening against CNBI-all, which solves every geometrically admissible subset. Nine synthetic scenarios are selected so that every level of decision dimension, dimensional excess, and dependence occurs three times. Each pair shares the exact payoff matrix and individual optima. This isolates the combination filter while retaining 10 seeds per scenario.")

add_heading(doc, "3.5 Statistical analysis", 2)
add_body(doc, "Method comparisons summarize the 10 seeds within each structural scenario and use the 27 scenarios as paired blocks. A Friedman test evaluates global differences, followed by paired Wilcoxon tests with Holm adjustment. Structural analysis of CNBI uses categorical decision dimension, dimensional excess, dependence, their interactions, and a seed block. Orthogonal contrasts are used to avoid artificial collinearity. Because the three-way model represents all 27 structural combinations, pure lack of fit cannot be separated within that saturated structural component.")

add_heading(doc, "4 Results", 1)
add_heading(doc, "4.1 Campaign validity", 2)
add_body(doc, "The controlled campaign completed 1,080 unique optimizer runs: 27 scenarios, 10 seeds, and four methods. No method produced an empty front. CNBI and rotated-factor NBI completed without evaluation ceilings, and all evolutionary runs respected their calibrated limits. The realized dependence targets remained within the declared tolerance.")

add_heading(doc, "4.2 Complete-front comparison", 2)
add_body(doc, "CNBI obtained the best median inverted generational distance in 26 of 27 scenarios and the best median hypervolume in 22. NSGA-III won the remaining inverted-generational-distance scenario and five hypervolume scenarios. CNBI therefore led the primary comparison in both coverage indicators. Its typical evaluation cost remained substantially below that of the evolutionary methods, although rotated-factor NBI was faster and cheaper at the cost of smaller, less comprehensive fronts.")
add_table(doc,
          ["Method", "Mean IGD rank", "Mean HV rank", "IGD wins", "HV wins", "Median evaluations"],
          [
              ("CNBI", "1.04", "1.22", "26/27", "22/27", "12,031"),
              ("NSGA-III", "1.96", "1.81", "1/27", "5/27", "84,000"),
              ("MOEA/D", "3.48", "3.04", "0/27", "0/27", "84,000"),
              ("Rotated-factor NBI", "3.52", "3.93", "0/27", "0/27", "1,095"),
          ], widths=[1.65, 1.0, 1.0, .85, .85, 1.15], caption="Table 2. Primary comparison using complete fronts")

add_heading(doc, "4.3 Sensitivity to output cardinality", 2)
add_body(doc, "The fixed-cardinality analysis included 270 paired blocks at 2 points, 260 at 5 points, 135 at 10 points, and 50 at 20 points. CNBI was 6.1% above the best median relative inverted generational distance with only 2 points and achieved the best median value from 5 through 20 points. Its relative hypervolume gap decreased from 62.4% at 2 points to 3.8% at 20 points. These curves show that CNBI benefits from retaining a moderately rich decision set; its advantage is reduced when every method must represent the frontier with only a few alternatives.")

add_heading(doc, "4.4 Structural effects", 2)
add_body(doc, "The factorial model explained 89.8% of the variation in inverted generational distance and 97.4% of the variation in hypervolume. Adjusted coefficients of determination were 88.3% and 97.0%, respectively. The three-way interaction among decision dimension, dimensional excess, and dependence was relevant for both responses. The effect of increasing any single factor therefore depends on the other two; no monotonic rule describes all structural conditions. The maximum variance inflation factor was 1 under orthogonal contrasts.")

add_heading(doc, "4.5 MaF stress tests", 2)
add_body(doc, "MaF8 and MaF9 produced stable CNBI approximations whose front sizes generally increased with objective count. Median inverted generational distance ranged from 0.0574 to 0.0780 for MaF8 and from 0.0699 to 0.1288 for MaF9. MaF13 was more difficult: median inverted generational distances were 0.3959, 0.4249, and 0.4767 for 8, 10, and 15 objectives. Final fronts contained approximately 17–19 points despite hundreds of accepted subproblem solutions.")
add_body(doc, "The MaF13 values were already normalized by the theoretical range of the true frontier. Median point-to-front distance remained between 0.065 and 0.076, but approximately 36.1% of points lay more than one normalized unit from the reference and the 95th percentile reached millions. A relatively small group of extreme solutions therefore dominated the mean generational distance. Repeated equations among objectives four onward also caused many objective combinations to collapse to the same decision vectors.")

add_heading(doc, "4.6 Turning application", 2)
add_body(doc, "CNBI and NSGA-III provided the strongest global coverage in the eight-response turning problem. CNBI achieved an inverted generational distance of 0.0754 and hypervolume of 0.5023. NSGA-III obtained 0.0958 and 0.5018, respectively. Rotated-factor NBI generated a sparser representation, whereas MOEA/D placed points close to a limited region of the reference frontier and consequently combined a low generational distance with weaker global coverage.")
add_table(doc,
          ["Method", "IGD", "GD", "HV"],
          [
              ("CNBI", "0.0754", "0.0269", "0.5023"),
              ("NSGA-III", "0.0958", "0.0282", "0.5018"),
              ("Rotated-factor NBI", "0.3782", "0.0375", "0.3236"),
              ("MOEA/D", "0.4257", "0.00073", "0.1979"),
          ], widths=[2.4, 1.25, 1.25, 1.25], caption="Table 3. Complete-front results for the turning application")

add_heading(doc, "4.7 Spectral-screening ablation", 2)
add_body(doc, "The median reduction in objective-vector evaluations ranged from 16.0% to 99.1% across the nine ablation scenarios, with a median scenario-level reduction of 70.5%. The largest saving occurred for nₓ = 5, m = 9, and high dependence. In that case, the median inverted-generational-distance increase was 0.0133 and the median hypervolume loss was 0.0232. Across the other eight scenarios, the inverted-generational-distance increase did not exceed 0.0067 and the hypervolume loss did not exceed 0.0111.")
add_table(doc,
          ["Scenario", "Evaluation reduction", "IGD increase", "HV loss"],
          [
              ("nₓ=2, m=4, low", "16.0%", "0.0034", "0.0049"),
              ("nₓ=2, m=6, medium", "25.7%", "0.0053", "0.0051"),
              ("nₓ=2, m=8, high", "72.6%", "0.0058", "0.0069"),
              ("nₓ=3, m=5, high", "94.7%", "0.0032", "0.0053"),
              ("nₓ=3, m=7, low", "46.1%", "0.0056", "0.0096"),
              ("nₓ=3, m=9, medium", "70.5%", "0.0067", "0.0072"),
              ("nₓ=5, m=7, medium", "96.7%", "0.0049", "0.0111"),
              ("nₓ=5, m=9, high", "99.1%", "0.0133", "0.0232"),
              ("nₓ=5, m=11, low", "63.8%", "0.0037", "0.0056"),
          ], widths=[2.2, 1.65, 1.2, 1.2], caption="Table 4. Paired ablation of spectral screening")

add_heading(doc, "5 Discussion", 1)
add_body(doc, "The results indicate that CNBI's primary advantage is coverage. Rather than committing the search to a single latent representation, the method assembles complementary information from several admissible objective subsets. Recomposition in the complete space preserves the original response meanings and allows all candidates to be compared under the same decision criteria. This behavior explains the strong complete-front inverted generational distance and hypervolume results.")
add_body(doc, "Cardinality is central to this interpretation. When all methods were compressed to two points, algorithms that retained a few extreme solutions were favored, particularly in hypervolume. As the fixed cardinality increased, CNBI approached the best hypervolume and led median coverage. Complete fronts and fixed-cardinality curves therefore answer different questions. The former measures the full service delivered by a method; the latter measures how effectively a restricted number of alternatives can be placed.")
add_body(doc, "The ablation supports the spectral mechanism. Screening removed little work in the smallest, weakly dependent scenario but removed almost all evaluations in some higher-dimensional structures. The associated quality loss was usually small. The largest loss occurred where the saving was also greatest, showing a transparent trade-off between exhaustive combinatorial search and computational economy. The results should not be reduced to a universal percentage because the benefit depends on the payoff geometry and the number of retained subsets.")
add_body(doc, "The turning case illustrates the decision-support value of a broad front. CNBI and NSGA-III achieved similar hypervolume, while CNBI obtained better inverted generational distance with fewer typical evaluations than the evolutionary methods in the synthetic campaign. Rotated-factor NBI remained computationally efficient but represented a smaller portion of the full response trade-off. MOEA/D's low generational distance demonstrates why proximity alone is insufficient: points may be close to the frontier while leaving large regions unexplored.")
add_body(doc, "MaF13 clarifies a limitation. Spectral structure alone cannot prevent different objective subsets from producing the same decision or from accepting numerically extreme candidates. Future versions should incorporate safeguards based on decision-space novelty, robust distance checks, or adaptive subset termination. These safeguards must preserve the method's defining property: final evaluation in the complete original objective space.")

add_heading(doc, "6 Limitations", 1)
add_body(doc, "The synthetic objectives are based on controlled squared-distance geometry, and other nonlinear landscapes may exhibit different behavior. The final MaF campaign evaluates CNBI only and is intended as a stress test rather than a formal comparison among methods. The nine-scenario ablation is balanced across factor levels but remains a selected subset of the 27 scenarios. The turning data were previously published and do not represent a new physical experiment. Finally, the extended multidimensional solver used here must be versioned separately from the historically registered CNBI 1.0 implementation.")

add_heading(doc, "7 Conclusions", 1)
add_body(doc, "CNBI extends NBI-style search to problems in which the number of objectives exceeds the conventional dimensional condition. It uses uncertainty-aware spectral screening to identify informative objective subsets and recomposes all solutions in the complete objective space. In the controlled synthetic campaign, CNBI led complete-front coverage in nearly every scenario and remained strong when cardinality was controlled. The nine-scenario ablation showed substantial, structure-dependent evaluation savings with limited quality loss. The turning application demonstrated practical competitiveness, while MaF13 exposed clear failure modes involving repetition and extreme solutions. Together, these findings support CNBI as an interpretable many-objective decision-support method and define specific directions for robust screening and adaptive cardinality in future work.")

add_heading(doc, "Data and code availability", 1)
add_body(doc, "The source code, experiment configurations, random seeds, raw outputs, and processed tables supporting this study will be deposited in a public repository. The archived repository link and DOI will be added before submission.", first_line=False)
add_heading(doc, "Declaration of generative AI and AI-assisted technologies in the manuscript preparation process", 1)
add_body(doc, "During the preparation of this work, the authors used OpenAI Codex to support manuscript organization, code review, data-analysis verification, and language editing. After using this tool, the authors reviewed and edited the content as needed and take full responsibility for the content of the published article.", first_line=False)

add_heading(doc, "References", 1)
references = [
    "Bao, C., Gao, D., Gu, W., Xu, L., & Goodman, E. D. (2023). A new adaptive decomposition-based evolutionary algorithm for multi- and many-objective optimization. Expert Systems with Applications, 213, Article 119080. https://doi.org/10.1016/j.eswa.2022.119080",
    "Cheng, R., Li, M., Tian, Y., Zhang, X., Yang, S., Jin, Y., & Yao, X. (2017). A benchmark test suite for evolutionary many-objective optimization. Complex & Intelligent Systems, 3, 67–81. https://doi.org/10.1007/s40747-017-0039-7",
    "Chiao, C.-H., & Hamada, M. (2001). Analyzing experiments with correlated multiple responses. Journal of Quality Technology, 33(4), 451–465. https://doi.org/10.1080/00224065.2001.11980104",
    "Das, I., & Dennis, J. E. (1998). Normal-boundary intersection: A new method for generating the Pareto surface in nonlinear multicriteria optimization problems. SIAM Journal on Optimization, 8(3), 631–657. https://doi.org/10.1137/S1052623496307510",
    "Deb, K., & Jain, H. (2014). An evolutionary many-objective optimization algorithm using reference-point-based nondominated sorting approach, Part I: Solving problems with box constraints. IEEE Transactions on Evolutionary Computation, 18(4), 577–601. https://doi.org/10.1109/TEVC.2013.2281535",
    "Horn, J. L. (1965). A rationale and test for the number of factors in factor analysis. Psychometrika, 30, 179–185. https://doi.org/10.1007/BF02289447",
    "Ishibuchi, H., Tsukamoto, N., & Nojima, Y. (2008). Evolutionary many-objective optimization: A short review. 2008 IEEE Congress on Evolutionary Computation, 2419–2426. https://doi.org/10.1109/CEC.2008.4631121",
    "Jolliffe, I. T., & Cadima, J. (2016). Principal component analysis: A review and recent developments. Philosophical Transactions of the Royal Society A: Mathematical, Physical and Engineering Sciences, 374(2065), Article 20150202. https://doi.org/10.1098/rsta.2015.0202",
    "Mechaacha, A., Belkaid, F., & Brahimi, N. (2026). Multi-objective multi-product process planning with reconfigurable machines: Exact and metaheuristic approaches. Expert Systems with Applications, 298, Article 129591. https://doi.org/10.1016/j.eswa.2025.129591",
    "Moura, D., Barcelos, V., Samanamud, G. R. L., França, A. B., Lofrano, R., Loures, C. C. A., Naves, L. L. R., Amaral, M. S., & Naves, F. L. (2018). Normal boundary intersection applied as multivariate and multiobjective optimization in the treatment of amoxicillin synthetic solution. Environmental Monitoring and Assessment, 190(3), Article 140. https://doi.org/10.1007/s10661-018-6523-8",
    "Myers, R. H., Montgomery, D. C., & Anderson-Cook, C. M. (2016). Response surface methodology: Process and product optimization using designed experiments (4th ed.). Wiley.",
    "Pereira, M. C., Ribeiro, C. T., Mendes, R. R. A., Campos, P. H. S., & Paiva, A. P. (2025). A hybrid multivariate normal boundary intersection approach with post-optimization assisted by mixture design of experiments. Engineering Applications of Artificial Intelligence, 162, Article 112510. https://doi.org/10.1016/j.engappai.2025.112510",
    "Purshouse, R. C., & Fleming, P. J. (2007). On the evolutionary optimization of many conflicting objectives. IEEE Transactions on Evolutionary Computation, 11(6), 770–784. https://doi.org/10.1109/TEVC.2007.910138",
    "Sinha, A., Saxena, D. K., Deb, K., & Tiwari, A. (2013). Using objective reduction and interactive procedure to handle many-objective optimization problems. Applied Soft Computing, 13(1), 415–427. https://doi.org/10.1016/j.asoc.2012.08.030",
    "Tong, L.-I., Wang, C.-H., & Chen, H.-C. (2005). Optimization of multiple responses using principal component analysis and technique for order preference by similarity to ideal solution. The International Journal of Advanced Manufacturing Technology, 27(3–4), 407–414. https://doi.org/10.1007/s00170-004-2157-9",
    "Zhang, Q., & Li, H. (2007). MOEA/D: A multiobjective evolutionary algorithm based on decomposition. IEEE Transactions on Evolutionary Computation, 11(6), 712–731. https://doi.org/10.1109/TEVC.2007.892759",
    "Zhang, W., Liu, J., Tan, S., & Wang, H. (2023). A decomposition-rotation dominance based evolutionary algorithm with reference point adaption for many-objective optimization. Expert Systems with Applications, 215, Article 119424. https://doi.org/10.1016/j.eswa.2022.119424",
]
for ref in references:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(ref)

props = doc.core_properties
props.title = "Combinatorial normal boundary intersection with spectral screening for structurally dependent many-objective optimization"
props.subject = "Initial manuscript draft for Expert Systems with Applications"
props.author = "Gabriel Victor de Lima; Mirelli de Castro Cesário; Matheus Costa Pereira; Anderson Paulo de Paiva"
props.keywords = "many-objective optimization, normal boundary intersection, spectral screening, parallel analysis"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print("Abstract words:", len(abstract.split()))
